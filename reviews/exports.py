import io
import re
from collections import Counter
from html import escape
from datetime import datetime, timezone

from django.db.models import Count
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

INK = colors.HexColor("#102d2b")
PAPER = colors.HexColor("#f4f0e8")
MUTED = colors.HexColor("#596e66")

EXCEL_HEADER_FILL = PatternFill("solid", fgColor="102D2B")
EXCEL_HEADER_FONT = Font(bold=True, color="F4F0E8", name="Calibri", size=11)
EXCEL_TITLE_FONT = Font(bold=True, color="102D2B", name="Calibri", size=16)
EXCEL_META_FONT = Font(color="596E66", name="Calibri", size=11)
EXCEL_ALT_FILL = PatternFill("solid", fgColor="F4F0E8")
EXCEL_BORDER = Border(
    left=Side(style="thin", color="D9DFD8"),
    right=Side(style="thin", color="D9DFD8"),
    top=Side(style="thin", color="D9DFD8"),
    bottom=Side(style="thin", color="D9DFD8"),
)


def _safe_filename(title: str) -> str:
    slug = re.sub(r"[^\w\s-]", "", title, flags=re.UNICODE).strip().lower()
    slug = re.sub(r"[-\s]+", "-", slug) or "app-reviews"
    return slug[:80]


def _sentiment_summary(reviews) -> dict[str, int]:
    counts = Counter(review.sentiment for review in reviews)
    return {
        "positive": counts.get("positive", 0),
        "negative": counts.get("negative", 0),
        "neutral": counts.get("neutral", 0),
    }


def build_excel_export(app, reviews) -> bytes:
    buffer = io.BytesIO()
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    sentiment = _sentiment_summary(reviews)

    summary["A1"] = app.title
    summary["A1"].font = EXCEL_TITLE_FONT
    summary["A2"] = app.developer or "—"
    summary["A2"].font = EXCEL_META_FONT
    summary["A3"] = app.app_id
    summary["A3"].font = EXCEL_META_FONT
    summary["A5"] = "Exported"
    summary["B5"] = generated
    summary["A6"] = "Total reviews"
    summary["B6"] = len(reviews)
    summary["A7"] = "Positive"
    summary["B7"] = sentiment["positive"]
    summary["A8"] = "Negative"
    summary["B8"] = sentiment["negative"]
    summary["A9"] = "Neutral"
    summary["B9"] = sentiment["neutral"]
    if app.store_score is not None:
        summary["A10"] = "Store rating"
        summary["B10"] = float(app.store_score)
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 42

    sheet = workbook.create_sheet("Reviews")
    headers = [
        "Rating",
        "Reviewer",
        "Sentiment",
        "Category",
        "Helpful",
        "Version",
        "Published",
        "Review",
        "AI insight",
        "Developer reply",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = EXCEL_HEADER_FILL
        cell.font = EXCEL_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = EXCEL_BORDER

    for index, review in enumerate(reviews, start=2):
        published = review.published_at.strftime("%Y-%m-%d %H:%M") if review.published_at else ""
        sheet.append(
            [
                review.rating,
                review.user_name or "Anonymous",
                review.sentiment.title(),
                review.category,
                review.thumbs_up,
                review.version or "",
                published,
                review.content,
                review.extracted_issue or "",
                review.reply_content or "",
            ]
        )
        for cell in sheet[index]:
            cell.border = EXCEL_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if index % 2 == 0:
                cell.fill = EXCEL_ALT_FILL

    widths = [8, 18, 12, 16, 9, 10, 18, 52, 36, 36]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"

    workbook.save(buffer)
    return buffer.getvalue()


def build_pdf_export(app, reviews) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=0.45 * inch,
        rightMargin=0.45 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        title=f"{app.title} — Reviews",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ExportTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        textColor=INK,
        spaceAfter=4,
    )
    meta_style = ParagraphStyle(
        "ExportMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        textColor=MUTED,
        spaceAfter=2,
    )
    cell_style = ParagraphStyle(
        "Cell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9,
        textColor=INK,
        alignment=TA_LEFT,
    )
    header_style = ParagraphStyle(
        "Header",
        parent=cell_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    sentiment = _sentiment_summary(reviews)
    story = [
        Paragraph(app.title, title_style),
        Paragraph(f"{app.developer or 'Unknown developer'} · {app.app_id}", meta_style),
        Paragraph(
            f"Exported {generated} · {len(reviews)} reviews · "
            f"+{sentiment['positive']} / −{sentiment['negative']} / ○{sentiment['neutral']}",
            meta_style,
        ),
        Spacer(1, 0.14 * inch),
    ]

    table_data = [
        [
            Paragraph("Rating", header_style),
            Paragraph("Reviewer", header_style),
            Paragraph("Sentiment", header_style),
            Paragraph("Category", header_style),
            Paragraph("Helpful", header_style),
            Paragraph("Published", header_style),
            Paragraph("Review & insight", header_style),
        ]
    ]
    for review in reviews:
        published = review.published_at.strftime("%Y-%m-%d") if review.published_at else "—"
        body = escape(review.content).replace("\n", "<br/>")
        if review.extracted_issue:
            body += f"<br/><font color='#e56b54'><b>Insight:</b> {escape(review.extracted_issue)}</font>"
        table_data.append(
            [
                Paragraph(str(review.rating), cell_style),
                Paragraph(review.user_name or "Anonymous", cell_style),
                Paragraph(review.sentiment.title(), cell_style),
                Paragraph(review.category, cell_style),
                Paragraph(str(review.thumbs_up), cell_style),
                Paragraph(published, cell_style),
                Paragraph(body, cell_style),
            ]
        )

    table = Table(
        table_data,
        colWidths=[0.55 * inch, 1.05 * inch, 0.75 * inch, 1.05 * inch, 0.6 * inch, 0.85 * inch, 5.55 * inch],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), INK),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PAPER]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d9dfd8")),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buffer.getvalue()


def export_filename(app, extension: str) -> str:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    return f"{_safe_filename(app.title)}-reviews-{stamp}.{extension}"
